import csv
import os
import sys
import sqlite3
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import date, datetime

APP_NAME = "نظام إدارة حضور ودرجات الطلاب"
TEACHER = "THAER ALJUMAILY"
APP_DIR = os.path.dirname(os.path.abspath(sys.executable if getattr(sys, "frozen", False) else __file__))
DB_FILE = os.path.join(APP_DIR, "school_system.db")
SECTIONS = ["A", "B", "C", "D", "E", "F", "G", "H"]
TERMS = ["الفصل الأول", "امتحان نصف السنة", "الفصل الثاني", "امتحان نهاية السنة"]
SCORE_FIELDS = [
    ("attendance", "الحضور", 10), ("absence", "الغياب", 10),
    ("participation", "المشاركة", 10), ("activity", "النشاط الصفي", 10),
    ("term1", "درجة امتحان الفصل الأول", 100),
    ("mid_year", "امتحان نصف السنة", 100), ("term2", "درجة امتحان الفصل الثاني", 100),
    ("end_year", "امتحان نهاية السنة", 100)
]

def db():
    return sqlite3.connect(DB_FILE)

def init_db():
    with db() as con:
        con.execute("""CREATE TABLE IF NOT EXISTS students(
            id INTEGER PRIMARY KEY AUTOINCREMENT, subject TEXT DEFAULT '',
            name TEXT NOT NULL, grade TEXT DEFAULT '', section TEXT DEFAULT '',
            attendance REAL DEFAULT 0, absence REAL DEFAULT 0,
            participation REAL DEFAULT 0, activity REAL DEFAULT 0,
            term1 REAL, mid_year REAL, term2 REAL, end_year REAL,
            term1_final REAL, term2_final REAL)""")
        columns = {row[1] for row in con.execute("PRAGMA table_info(students)")}
        for column in ("term1_final", "term2_final"):
            if column not in columns:
                con.execute(f"ALTER TABLE students ADD COLUMN {column} REAL")
        # Migrate the former shared values once as starting values for term 2,
        # preserving existing totals; after migration the term columns are independent.
        legacy_to_term2 = {
            "attendance2": "attendance", "absence2": "absence",
            "participation2": "participation", "activity2": "activity",
        }
        for new_column, old_column in legacy_to_term2.items():
            if new_column not in columns:
                con.execute(f"ALTER TABLE students ADD COLUMN {new_column} REAL")
                con.execute(f"UPDATE students SET {new_column}={old_column} WHERE {new_column} IS NULL")
        con.execute("""CREATE TABLE IF NOT EXISTS participations(
            id INTEGER PRIMARY KEY AUTOINCREMENT, student_id INTEGER,
            student_name TEXT NOT NULL, grade TEXT DEFAULT '', section TEXT DEFAULT '',
            subject TEXT DEFAULT '', term TEXT NOT NULL, lesson_date TEXT NOT NULL,
            activity_score REAL DEFAULT 0, notes TEXT DEFAULT '')""")
        con.execute("""CREATE TABLE IF NOT EXISTS settings(
            key TEXT PRIMARY KEY, value TEXT NOT NULL)""")
        for key, _, maximum in SCORE_FIELDS:
            con.execute("INSERT OR IGNORE INTO settings VALUES(?,?)",
                        ("max_" + key, str(maximum)))
        con.execute("INSERT OR IGNORE INTO settings VALUES('subject','')")
        con.execute("INSERT OR IGNORE INTO settings VALUES('term',?)", (TERMS[0],))

def setting(key, default=""):
    with db() as con:
        row = con.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    return row[0] if row else default

def save_setting(key, value):
    with db() as con:
        con.execute("""INSERT INTO settings(key,value) VALUES(?,?)
            ON CONFLICT(key) DO UPDATE SET value=excluded.value""", (key, str(value)))

def num(entry, label, blank=None):
    raw = entry.get().strip()
    if not raw:
        return blank
    try:
        value = float(raw)
    except ValueError:
        raise ValueError(f"حقل «{label}» يجب أن يحتوي على رقم.")
    if value < 0:
        raise ValueError(f"لا يمكن أن تكون «{label}» سالبة.")
    return value

def fmt(value):
    return "غير مدخل" if value is None else f"{value:g}"

def average(values):
    # Empty legacy fields were imported as 0 in older versions; do not count those as grades.
    entered = [v for v in values if v is not None and v != 0]
    return None if not entered else round(sum(entered) / len(entered), 2)

def term1_total(row):
    return None if row[9] is None else (row[5] or 0) + (row[6] or 0) + (row[7] or 0) + (row[8] or 0) + row[9]

def term2_total(row):
    # Second-term components are stored separately; the same four form controls edit them by term selection.
    return None if row[11] is None else (row[15] or 0) + (row[16] or 0) + (row[17] or 0) + (row[18] or 0) + row[11]

def final_average(row):
    # Average only grades actually entered: calculated term totals, mid-year, and year-end.
    return average([term1_total(row), row[10], term2_total(row), row[12]])

def excel_export_value(value):
    return value if value is not None else "غير مدخل"


class StudentApp:
    def __init__(self, root):
        self.root = root
        self.root.title(f"{APP_NAME} - {TEACHER}")
        self.root.geometry("1500x900")
        self.root.minsize(1100, 680)
        self.root.configure(bg="#eef2f7")
        self.selected_id = None
        self.component_cache = {"first": [0, 0, 0, 0], "second": [0, 0, 0, 0]}
        self.component_term_key = "first"
        init_db()
        self.maxima = {k: float(setting("max_" + k, m)) for k, _, m in SCORE_FIELDS}
        self.make_header()
        self.make_form()
        self.make_buttons()
        self.make_filters()
        self.make_tables()
        self.refresh()
        self.root.protocol("WM_DELETE_WINDOW", self.close)

    def make_header(self):
        f = tk.Frame(self.root, bg="#182957", pady=10)
        f.pack(fill="x")
        tk.Label(f, text=APP_NAME, font=("Arial", 22, "bold"),
                 bg="#182957", fg="white").pack()
        tk.Label(f, text=TEACHER, font=("Arial", 12, "bold"),
                 bg="#182957", fg="#d9e5ff").pack()

    def field(self, parent, label, width=15):
        f = tk.Frame(parent, bg="white")
        tk.Label(f, text=label, bg="white", fg="#24324b",
                 font=("Arial", 10, "bold")).pack(anchor="e")
        e = tk.Entry(f, width=width, font=("Arial", 12), justify="right")
        e.pack(fill="x", pady=(3, 0))
        return f, e

    def make_form(self):
        box = tk.LabelFrame(self.root, text="بيانات الطالب والدرجات",
                            font=("Arial", 12, "bold"), bg="white",
                            fg="#182957", padx=10, pady=8)
        box.pack(fill="x", padx=12, pady=8)
        for c in range(5):
            box.grid_columnconfigure(c, weight=1)

        self.name = self.grade = None
        self.attendance = self.absence = self.participation = self.activity = None
        self.term1 = self.mid_year = self.term2 = self.end_year = None
        self.section = self.subject = self.term = None

        def combo_field(label, values, attr, readonly=False):
            frame = tk.Frame(box, bg="white")
            tk.Label(frame, text=label, bg="white", fg="#24324b",
                     font=("Arial", 10, "bold")).pack(anchor="e")
            widget = ttk.Combobox(frame, values=values, state="readonly" if readonly else "normal",
                                  justify="right", font=("Arial", 11))
            widget.pack(fill="x", pady=(3, 0))
            setattr(self, attr, widget)
            return frame, widget

        first_row = [
            ("اسم الطالب", "name", 22, "entry"), ("الصف", "grade", 12, "grade"),
            ("الشعبة", "section", 10, "section"), ("اسم المادة الدراسية", "subject", 18, "subject"),
            ("الفصل الدراسي", "term", 14, "term"),
        ]
        for col, (label, attr, width, kind) in enumerate(first_row):
            if kind == "entry":
                frame, widget = self.field(box, label, width)
                setattr(self, attr, widget)
            elif kind == "grade":
                frame, widget = combo_field(label, [], attr)
            elif kind == "section":
                frame, widget = combo_field(label, SECTIONS, attr, readonly=True)
                widget.set(SECTIONS[0])
            elif kind == "subject":
                frame = tk.Frame(box, bg="white")
                tk.Label(frame, text=label, bg="white", fg="#24324b",
                         font=("Arial", 10, "bold")).pack(anchor="e")
                widget = tk.Entry(frame, width=width, font=("Arial", 12), justify="right")
                widget.insert(0, setting("subject")); widget.pack(fill="x", pady=(3, 0))
                self.subject = widget
            else:
                frame, widget = combo_field(label, TERMS, attr, readonly=True)
                widget.set(setting("term", TERMS[0]))
                widget.bind("<<ComboboxSelected>>", self.on_term_changed)
            frame.grid(row=0, column=4-col, sticky="ew", padx=4, pady=4)

        def add_score_row(row, items):
            for col, (label, attr, width) in enumerate(items):
                frame, widget = self.field(box, label, width)
                setattr(self, attr, widget)
                frame.grid(row=row, column=4-col, sticky="ew", padx=4, pady=4)

        add_score_row(1, [
            ("الحضور", "attendance", 10), ("الغياب", "absence", 10),
            ("المشاركة", "participation", 10), ("النشاط الصفي", "activity", 10),
        ])
        add_score_row(2, [
            ("امتحان الفصل الأول", "term1", 10), ("امتحان نصف السنة", "mid_year", 10),
            ("امتحان الفصل الثاني", "term2", 10), ("امتحان نهاية السنة", "end_year", 10),
        ])
        self.total_label = tk.Label(box, text="المعدل النهائي للدرجات المدخلة: غير مكتمل",
                                    bg="white", fg="#135ec9", font=("Arial", 12, "bold"))
        self.total_label.grid(row=3, column=0, columnspan=5, sticky="e", padx=5, pady=6)
        for entry in (self.attendance, self.absence, self.participation, self.activity,
                      self.term1, self.mid_year, self.term2, self.end_year):
            entry.bind("<KeyRelease>", self.update_total)

    def _component_key(self):
        return "second" if self.term.get() == TERMS[2] else "first"

    def _read_component_entries(self):
        return [
            num(self.attendance, "الحضور", 0) or 0,
            num(self.absence, "الغياب", 0) or 0,
            num(self.participation, "المشاركة", 0) or 0,
            num(self.activity, "النشاط الصفي", 0) or 0,
        ]

    def _show_component_values(self, values):
        for entry, value in zip((self.attendance, self.absence, self.participation, self.activity), values):
            entry.delete(0, tk.END)
            if value is not None:
                entry.insert(0, str(value))

    def _capture_component_values(self):
        key = self._component_key()
        self.component_cache[key] = self._read_component_entries()
        self.component_term_key = key

    def on_term_changed(self, event=None):
        # Save the visible set in memory before displaying the chosen semester's set.
        previous = getattr(self, "component_term_key", self._component_key())
        try:
            self.component_cache[previous] = self._read_component_entries()
        except ValueError:
            pass
        save_setting("term", self.term.get())
        chosen = self._component_key()
        self.component_term_key = chosen
        self._show_component_values(self.component_cache[chosen])
        self.update_total()

    def make_buttons(self):
        f = tk.Frame(self.root, bg="#eef2f7")
        f.pack(fill="x", padx=12, pady=5)
        buttons = [
            ("إضافة طالب", "#198754", self.add_student),
            ("حفظ التعديل", "#0d6efd", self.update_student),
            ("مسح الحقول", "#6c757d", self.clear),
            ("حذف الطالب", "#dc3545", self.delete_student),
            ("تسجيل المشاركين", "#7c3aed", self.participation_window),
            ("إعدادات الدرجات", "#b8860b", self.settings_window),
            ("إصدار", "#ffbf69", self.export_csv),
            ("إدراج", "#75d6c4", self.import_csv),
            ("طباعة", "#334155", self.print_report),
            ("نسخ احتياطي", "#475569", self.backup),
        ]
        right_group = tk.Frame(f, bg="#eef2f7")
        right_group.pack(side="right")
        left_group = tk.Frame(f, bg="#eef2f7")
        left_group.pack(side="left")
        for index, (label, color, command) in enumerate(buttons):
            parent = right_group if index < 4 else left_group
            side = "right" if index < 4 else "left"
            tk.Button(parent, text=label, command=command, bg=color, fg="white",
                      font=("Arial", 10, "bold"), relief="flat",
                      padx=7, pady=8).pack(side=side, padx=3, pady=3)

    def make_filters(self):
        f = tk.LabelFrame(self.root, text="البحث والتصفية", bg="#eef2f7",
                          fg="#182957", font=("Arial", 10, "bold"))
        f.pack(fill="x", padx=12, pady=4)
        tk.Label(f, text="اسم الطالب:", bg="#eef2f7").pack(side="right", padx=4)
        self.search = tk.StringVar()
        e = tk.Entry(f, textvariable=self.search, justify="right", width=20)
        e.pack(side="right", padx=4)
        e.bind("<KeyRelease>", lambda ev: self.refresh())
        tk.Label(f, text="الشعبة:", bg="#eef2f7").pack(side="right", padx=4)
        self.filter_section = ttk.Combobox(f, values=["الكل"] + SECTIONS,
                                           state="readonly", width=9)
        self.filter_section.set("الكل")
        self.filter_section.pack(side="right", padx=4)
        self.filter_section.bind("<<ComboboxSelected>>", lambda ev: self.refresh())
        tk.Label(f, text="الصف:", bg="#eef2f7").pack(side="right", padx=4)
        self.filter_grade = ttk.Combobox(f, values=["الكل"], state="readonly", width=12)
        self.filter_grade.set("الكل")
        self.filter_grade.pack(side="right", padx=4)
        self.filter_grade.bind("<<ComboboxSelected>>", lambda ev: self.refresh())
        self.count_label = tk.Label(f, text="عدد الطلاب: 0", bg="#eef2f7",
                                    fg="#135ec9", font=("Arial", 10, "bold"))
        self.count_label.pack(side="left", padx=8)

    def make_tree(self, parent, columns, labels):
        frame = tk.Frame(parent)
        frame.pack(fill="both", expand=True)
        y = ttk.Scrollbar(frame, orient="vertical")
        x = ttk.Scrollbar(frame, orient="horizontal")
        tree = ttk.Treeview(frame, columns=columns, show="headings",
                            yscrollcommand=y.set, xscrollcommand=x.set)
        y.config(command=tree.yview)
        x.config(command=tree.xview)
        for col in columns:
            tree.heading(col, text=labels[col], anchor="center")
            tree.column(col, width=150 if col in ("name", "student_name") else 115,
                        anchor="center", stretch=True)
        tree.grid(row=0, column=0, sticky="nsew")
        y.grid(row=0, column=1, sticky="ns")
        x.grid(row=1, column=0, sticky="ew")
        frame.grid_rowconfigure(0, weight=1)
        frame.grid_columnconfigure(0, weight=1)
        return tree

    def make_tables(self):
        style = ttk.Style()
        style.configure("Treeview", font=("Arial", 10), rowheight=28)
        style.configure("Treeview.Heading", font=("Arial", 10, "bold"))
        style.configure("TNotebook", background="#eef2f7", borderwidth=0)
        style.configure("TNotebook.Tab", background="#dce8f8", foreground="#182957",
                        padding=(12, 7), font=("Arial", 10, "bold"))
        style.map("TNotebook.Tab",
                  background=[("selected", "#ffffff"), ("active", "#edf4ff")],
                  foreground=[("selected", "#135ec9"), ("active", "#182957")])
        book = ttk.Notebook(self.root)
        book.pack(fill="both", expand=True, padx=12, pady=5)
        first, final, participation, second, midyear, endyear = [tk.Frame(book) for _ in range(6)]
        book.add(first, text="درجات الفصل الأول")
        book.add(second, text="درجات الفصل الثاني")
        book.add(midyear, text="نصف السنة")
        book.add(endyear, text="نهاية السنة")
        book.add(final, text="الدرجات النهائية")
        book.add(participation, text="سجل المشاركات")
        # Tkinter Treeview displays declared columns left-to-right; reverse declaration
        # so the visible order reads from the right: ت، اسم الطالب، الصف، الشعبة...
        self.daily_cols = ("term1_total", "term1", "activity", "participation", "absence", "attendance",
                           "section", "grade", "name", "id")
        self.daily_labels = {"id":"ت", "name":"اسم الطالب", "grade":"الصف",
            "section":"الشعبة", "attendance":"الحضور", "absence":"الغياب",
            "participation":"المشاركة", "activity":"النشاط الصفي",
            "term1":"امتحان الفصل الأول", "term1_total":"مجموع الفصل الأول"}
        self.daily_tree = self.make_tree(first, self.daily_cols, self.daily_labels)
        self.daily_tree.bind("<<TreeviewSelect>>", self.select_student)

        self.term2_cols = ("term2_total", "term2", "activity", "participation", "absence", "attendance",
                           "section", "grade", "name", "id")
        self.term2_labels = {"id":"ت", "name":"اسم الطالب", "grade":"الصف", "section":"الشعبة",
            "attendance":"الحضور", "absence":"الغياب", "participation":"المشاركة",
            "activity":"النشاط الصفي", "term2":"امتحان الفصل الثاني", "term2_total":"مجموع الفصل الثاني"}
        self.term2_tree = self.make_tree(second, self.term2_cols, self.term2_labels)
        self.term2_tree.bind("<<TreeviewSelect>>", self.select_student)

        identity = ("section", "grade", "name", "id")
        self.midyear_cols = ("mid_year",) + identity
        self.midyear_labels = {"id":"ت", "name":"اسم الطالب", "grade":"الصف", "section":"الشعبة",
                               "mid_year":"درجة نصف السنة"}
        self.midyear_tree = self.make_tree(midyear, self.midyear_cols, self.midyear_labels)
        self.midyear_tree.bind("<<TreeviewSelect>>", self.select_student)
        self.endyear_cols = ("end_year",) + identity
        self.endyear_labels = {"id":"ت", "name":"اسم الطالب", "grade":"الصف", "section":"الشعبة",
                               "end_year":"درجة نهاية السنة"}
        self.endyear_tree = self.make_tree(endyear, self.endyear_cols, self.endyear_labels)
        self.endyear_tree.bind("<<TreeviewSelect>>", self.select_student)

        self.final_cols = ("average", "end_year", "term2_total", "mid_year", "term1_total",
                           "section", "grade", "name", "id")
        self.final_labels = {"id":"ت", "name":"اسم الطالب", "grade":"الصف", "section":"الشعبة",
            "term1_total":"درجة الفصل الأول", "mid_year":"درجة نصف السنة",
            "term2_total":"درجة الفصل الثاني", "end_year":"درجة نهاية السنة",
            "average":"المعدل النهائي"}
        self.final_tree = self.make_tree(final, self.final_cols, self.final_labels)
        self.final_tree.bind("<<TreeviewSelect>>", self.select_student)

        self.part_cols = ("activity_score", "lesson_date", "term", "subject",
                          "section", "grade", "student_name", "id")
        self.part_labels = {"id":"ت", "student_name":"اسم الطالب", "grade":"الصف",
            "section":"الشعبة", "subject":"المادة", "term":"الفصل الدراسي",
            "lesson_date":"تاريخ الدرس", "activity_score":"درجة النشاط الصفي"}
        self.part_tree = self.make_tree(participation, self.part_cols, self.part_labels)

    def read_form(self):
        name = self.name.get().strip()
        if not name:
            raise ValueError("يرجى كتابة اسم الطالب.")
        self._capture_component_values()
        first, second = self.component_cache["first"], self.component_cache["second"]
        data = {
            "subject": self.subject.get().strip(), "name": name,
            "grade": self.grade.get().strip(), "section": self.section.get().strip(),
            "attendance": first[0], "absence": first[1], "participation": first[2], "activity": first[3],
            "attendance2": second[0], "absence2": second[1],
            "participation2": second[2], "activity2": second[3],
            "term1": num(self.term1, "درجة امتحان الفصل الأول"),
            "mid_year": num(self.mid_year, "درجة امتحان نصف السنة"),
            "term2": num(self.term2, "درجة امتحان الفصل الثاني"),
            "end_year": num(self.end_year, "درجة امتحان نهاية السنة")
        }
        if not data["section"]:
            raise ValueError("اختر الشعبة.")
        for key, label, _ in SCORE_FIELDS:
            v = data[key]
            if v is not None and v > self.maxima[key]:
                raise ValueError(f"درجة «{label}» أكبر من الحد الأعلى ({self.maxima[key]:g}).")
        for key, base, label in (("attendance2","attendance","الحضور - الفصل الثاني"),
                                 ("absence2","absence","الغياب - الفصل الثاني"),
                                 ("participation2","participation","المشاركة - الفصل الثاني"),
                                 ("activity2","activity","النشاط الصفي - الفصل الثاني")):
            if data[key] > self.maxima[base]:
                raise ValueError(f"درجة «{label}» أكبر من الحد الأعلى ({self.maxima[base]:g}).")
        return data

    def update_total(self, event=None):
        try:
            self._capture_component_values()
            first_values = self.component_cache["first"]
            second_values = self.component_cache["second"]
            first_exam = num(self.term1, "درجة امتحان الفصل الأول", None)
            mid_year = num(self.mid_year, "درجة امتحان نصف السنة", None)
            second_exam = num(self.term2, "درجة امتحان الفصل الثاني", None)
            end_year = num(self.end_year, "درجة امتحان نهاية السنة", None)
            first_total = None if first_exam is None else sum(first_values) + first_exam
            second_total = None if second_exam is None else sum(second_values) + second_exam
            result = average([first_total, mid_year, second_total, end_year])
            self.total_label.config(
                text=f"المعدل النهائي للدرجات المدخلة: {result:.2f} / 100" if result is not None
                else "المعدل النهائي للدرجات المدخلة: غير مكتمل")
        except ValueError:
            self.total_label.config(text="المعدل النهائي للدرجات المدخلة: أكمل الدرجات بصيغة رقمية")

    def add_student(self):
        try:
            d = self.read_form()
            with db() as con:
                con.execute("""INSERT INTO students
                    (subject,name,grade,section,attendance,absence,participation,activity,
                     term1,mid_year,term2,end_year,attendance2,absence2,participation2,activity2)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    tuple(d[k] for k in ("subject","name","grade","section","attendance",
                    "absence","participation","activity","term1","mid_year","term2","end_year",
                    "attendance2","absence2","participation2","activity2")))
            save_setting("subject", d["subject"])
            self.clear()
            self.refresh()
            messagebox.showinfo("نجاح", "تمت إضافة الطالب.")
        except (ValueError, sqlite3.Error) as e:
            messagebox.showerror("تعذر الإضافة", str(e))

    def update_student(self):
        if self.selected_id is None:
            messagebox.showwarning("تنبيه", "اختر طالبًا من الجدول أولًا.")
            return
        try:
            d = self.read_form()
            with db() as con:
                con.execute("""UPDATE students SET subject=?,name=?,grade=?,section=?,
                    attendance=?,absence=?,participation=?,activity=?,term1=?,mid_year=?,
                    term2=?,end_year=?,attendance2=?,absence2=?,participation2=?,activity2=?,
                    term1_final=NULL,term2_final=NULL WHERE id=?""",
                    tuple(d[k] for k in ("subject","name","grade","section","attendance",
                    "absence","participation","activity","term1","mid_year","term2","end_year",
                    "attendance2","absence2","participation2","activity2")) + (self.selected_id,))
            self.clear()
            self.refresh()
            messagebox.showinfo("نجاح", "تم تحديث بيانات الطالب.")
        except (ValueError, sqlite3.Error) as e:
            messagebox.showerror("تعذر التعديل", str(e))

    def delete_student(self):
        if self.selected_id is None:
            messagebox.showwarning("تنبيه", "اختر طالبًا من الجدول أولًا.")
            return
        if not messagebox.askyesno("تأكيد", "هل تريد حذف الطالب المحدد؟"):
            return
        with db() as con:
            con.execute("DELETE FROM students WHERE id=?", (self.selected_id,))
        self.clear()
        self.refresh()

    def clear(self):
        self.selected_id = None
        self.component_cache = {"first": [0, 0, 0, 0], "second": [0, 0, 0, 0]}
        self.component_term_key = self._component_key()
        for e in (self.name, self.grade, self.attendance, self.absence, self.participation,
                  self.activity, self.term1, self.mid_year, self.term2, self.end_year):
            e.delete(0, tk.END)
        self.section.set(SECTIONS[0])
        self.total_label.config(text="المعدل النهائي للدرجات المدخلة: غير مكتمل")

    def select_student(self, event):
        tree = event.widget
        selected = tree.selection()
        if not selected:
            return
        try:
            sid = int(selected[0])
        except ValueError:
            return
        with db() as con:
            r = con.execute("SELECT * FROM students WHERE id=?", (sid,)).fetchone()
        if not r:
            return
        self.selected_id = r[0]
        self.clear()
        self.selected_id = r[0]
        self.subject.delete(0, tk.END); self.subject.insert(0, r[1] or "")
        self.name.insert(0, r[2] or "")
        self.grade.insert(0, r[3] or "")
        self.section.set(r[4] if r[4] in SECTIONS else SECTIONS[0])
        self.component_cache = {
            "first": [r[5] or 0, r[6] or 0, r[7] or 0, r[8] or 0],
            "second": [r[15] or 0, r[16] or 0, r[17] or 0, r[18] or 0],
        }
        self.component_term_key = self._component_key()
        self._show_component_values(self.component_cache[self.component_term_key])
        for e, v in zip((self.term1,self.mid_year,self.term2,self.end_year), r[9:13]):
            if v is not None:
                e.insert(0, str(v))
        self.update_total()

    def refresh(self):
        if not hasattr(self, "daily_tree"):
            return
        for tree in (self.daily_tree, self.term2_tree, self.midyear_tree, self.endyear_tree, self.final_tree, self.part_tree):
            for item in tree.get_children():
                tree.delete(item)
        where, params = [], []
        if hasattr(self, "filter_grade"):
            grade_filter = self.filter_grade.get()
            if grade_filter and grade_filter != "الكل":
                where.append("grade=?"); params.append(grade_filter)
        if self.search.get().strip():
            where.append("name LIKE ?"); params.append("%" + self.search.get().strip() + "%")
        sec = self.filter_section.get() if self.filter_section.get() else "الكل"
        if sec != "الكل":
            where.append("section=?"); params.append(sec)
        query = "SELECT * FROM students" + ((" WHERE " + " AND ".join(where)) if where else "")
        query += " ORDER BY section, name COLLATE NOCASE"
        with db() as con:
            students = con.execute(query, params).fetchall()
            parts = con.execute("""SELECT id,student_name,grade,section,subject,term,
                lesson_date,activity_score FROM participations ORDER BY lesson_date DESC,id DESC""").fetchall()
        with db() as con:
            grades = [r[0] for r in con.execute("SELECT DISTINCT grade FROM students WHERE grade IS NOT NULL AND grade != '' ORDER BY grade").fetchall()]
        if hasattr(self, "filter_grade"):
            current_grade = self.filter_grade.get()
            self.filter_grade.configure(values=["الكل"] + grades)
            if current_grade not in ["الكل"] + grades: self.filter_grade.set("الكل")
        self.grade.configure(values=grades)
        self.count_label.config(text=f"عدد الطلاب: {len(students)}")
        section_serials = {}
        for r in students:
            section_key = (r[4] or "غير محددة").strip() or "غير محددة"
            section_serials[section_key] = section_serials.get(section_key, 0) + 1
            display_serial = section_serials[section_key]
            avg = final_average(r)
            # Final first-term total = exam + attendance + participation + absence + classroom activity.
            # Absence is added as an entered score, not subtracted. Do not show a total until the exam is entered.
            term1_total = None if r[9] is None else round((r[5] or 0) + (r[6] or 0) + (r[7] or 0) + (r[8] or 0) + r[9], 2)
            daily = {"id":display_serial, "name":r[2], "grade":r[3], "section":r[4],
                "attendance":fmt(r[5]), "absence":fmt(r[6]), "participation":fmt(r[7]),
                "activity":fmt(r[8]), "term1":fmt(r[9]), "term1_total":fmt(term1_total)}
            self.daily_tree.insert("", "end", iid=str(r[0]), values=tuple(daily[c] for c in self.daily_cols))
            term2_total_value = term2_total(r)
            self.midyear_tree.insert("", "end", iid=str(r[0]), values=(fmt(r[10]), r[4], r[3], r[2], display_serial))
            self.endyear_tree.insert("", "end", iid=str(r[0]), values=(fmt(r[12]), r[4], r[3], r[2], display_serial))
            final_values = {"id":display_serial, "name":r[2], "grade":r[3], "section":r[4],
                "term1_total":fmt(term1_total), "mid_year":fmt(r[10]),
                "term2_total":fmt(term2_total_value),
                "end_year":fmt(r[12]), "average":fmt(avg)}
            self.final_tree.insert("", "end", iid=str(r[0]), values=tuple(final_values[c] for c in self.final_cols))
            term2_values = {"id":display_serial, "name":r[2], "grade":r[3], "section":r[4],
                "attendance":fmt(r[15]), "absence":fmt(r[16]), "participation":fmt(r[17]),
                "activity":fmt(r[18]), "term2":fmt(r[11]), "term2_total":fmt(term2_total_value)}
            self.term2_tree.insert("", "end", iid=str(r[0]), values=tuple(term2_values[c] for c in self.term2_cols))
        for r in parts:
            d = dict(zip(("id","student_name","grade","section","subject","term","lesson_date","activity_score"),r))
            self.part_tree.insert("", "end", iid="p"+str(r[0]),
                                  values=tuple(d[c] for c in self.part_cols))

    def participation_window(self):
        win = tk.Toplevel(self.root)
        win.title("تسجيل المشاركين - النشاط الصفي")
        win.geometry("560x500")
        win.configure(bg="white")
        labels = ["اسم الطالب", "الصف", "اسم المادة الدراسية", "درجة النشاط الصفي", "ملاحظات"]
        entries = {}
        for i, label in enumerate(labels):
            tk.Label(win, text=label, bg="white", font=("Arial", 10, "bold")).grid(
                row=i, column=1, sticky="e", padx=12, pady=6)
            e = tk.Entry(win, font=("Arial", 12), justify="right", width=32)
            e.grid(row=i, column=0, sticky="ew", padx=12, pady=6)
            entries[label] = e
        tk.Label(win, text="الشعبة", bg="white").grid(row=5,column=1,sticky="e",padx=12,pady=6)
        section = ttk.Combobox(win, values=SECTIONS, state="readonly", justify="center")
        section.set(self.section.get() or SECTIONS[0]); section.grid(row=5,column=0,sticky="ew",padx=12,pady=6)
        tk.Label(win, text="الفصل الدراسي", bg="white").grid(row=6,column=1,sticky="e",padx=12,pady=6)
        term = ttk.Combobox(win, values=TERMS, state="readonly", justify="right")
        term.set(self.term.get()); term.grid(row=6,column=0,sticky="ew",padx=12,pady=6)
        tk.Label(win, text="تاريخ الدرس YYYY-MM-DD", bg="white").grid(row=7,column=1,sticky="e",padx=12,pady=6)
        date_entry = tk.Entry(win, justify="right")
        date_entry.insert(0, date.today().isoformat()); date_entry.grid(row=7,column=0,sticky="ew",padx=12,pady=6)
        if self.selected_id:
            with db() as con:
                r = con.execute("SELECT name,grade,section,subject FROM students WHERE id=?",
                                (self.selected_id,)).fetchone()
            if r:
                for label, value in zip(labels[:3], r[:2] + (r[3],)):
                    entries[label].insert(0, value or "")
                section.set(r[2] if r[2] in SECTIONS else SECTIONS[0])

        def save():
            try:
                name = entries["اسم الطالب"].get().strip()
                grade = entries["الصف"].get().strip()
                subject = entries["اسم المادة الدراسية"].get().strip()
                notes = entries["ملاحظات"].get().strip()
                if not name: raise ValueError("اكتب اسم الطالب المشارك.")
                lesson_date = date_entry.get().strip()
                datetime.strptime(lesson_date, "%Y-%m-%d")
                score = num(entries["درجة النشاط الصفي"], "درجة النشاط الصفي", 0)
                if score > self.maxima["activity"]:
                    raise ValueError(f"الدرجة تتجاوز الحد الأعلى للنشاط الصفي ({self.maxima['activity']:g}).")
                with db() as con:
                    duplicate = con.execute("""SELECT id FROM participations WHERE
                        student_name=? AND grade=? AND section=? AND subject=? AND term=? AND lesson_date=?""",
                        (name,grade,section.get(),subject,term.get(),lesson_date)).fetchone()
                    if duplicate: raise ValueError("هذه المشاركة مسجلة مسبقًا لهذا الطالب في التاريخ نفسه.")
                    student = con.execute("""SELECT id FROM students WHERE name=? AND grade=? AND section=?
                        ORDER BY id LIMIT 1""", (name,grade,section.get())).fetchone()
                    con.execute("""INSERT INTO participations(student_id,student_name,grade,section,
                        subject,term,lesson_date,activity_score,notes) VALUES(?,?,?,?,?,?,?,?,?)""",
                        (student[0] if student else None,name,grade,section.get(),subject,term.get(),
                         lesson_date,score,notes))
                self.refresh(); win.destroy()
                messagebox.showinfo("نجاح", "تم تسجيل المشاركة.")
            except (ValueError, sqlite3.Error) as ex:
                messagebox.showerror("تعذر التسجيل", str(ex), parent=win)
        tk.Button(win, text="حفظ المشاركة", command=save, bg="#7c3aed", fg="white",
                  font=("Arial",11,"bold"), pady=8).grid(row=8,column=0,columnspan=2,pady=12)
        win.grid_columnconfigure(0, weight=1)

    def settings_window(self):
        win = tk.Toplevel(self.root)
        win.title("إعدادات الحدود العليا للدرجات")
        win.geometry("480x430")
        win.configure(bg="white")
        entries = {}
        tk.Label(win, text="اضبط الحد الأعلى لكل حقل وفق نظام مدرستك",
                 bg="white", font=("Arial", 11, "bold")).pack(pady=10)
        for key, label, default in SCORE_FIELDS:
            row = tk.Frame(win, bg="white"); row.pack(fill="x", padx=12, pady=4)
            tk.Label(row, text=label, bg="white", width=25, anchor="e").pack(side="right", padx=6)
            e = tk.Entry(row, width=12, justify="center")
            e.insert(0, str(self.maxima[key])); e.pack(side="right")
            entries[key] = e
        def save():
            try:
                vals = {}
                for key, e in entries.items():
                    value = float(e.get())
                    if value <= 0: raise ValueError("يجب أن يكون الحد الأعلى أكبر من صفر.")
                    vals[key] = value
                for key, value in vals.items(): save_setting("max_" + key, value)
                self.maxima = vals
                win.destroy()
                messagebox.showinfo("نجاح", "تم حفظ الحدود العليا.")
            except ValueError as ex: messagebox.showerror("خطأ", str(ex), parent=win)
        tk.Button(win, text="حفظ", command=save, bg="#198754", fg="white",
                  font=("Arial",11,"bold"), padx=15,pady=7).pack(pady=12)

    def export_csv(self):
        path = filedialog.asksaveasfilename(defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")], title="تصدير Excel - ورقة مستقلة لكل شعبة")
        if not path: return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter

            with db() as con:
                rows = con.execute("SELECT * FROM students ORDER BY section,name COLLATE NOCASE").fetchall()
            if not rows:
                messagebox.showwarning("تنبيه", "لا توجد بيانات للتصدير.")
                return

            groups = {}
            for row in rows:
                section = (row[4] or "غير محددة").strip() or "غير محددة"
                groups.setdefault(section, []).append(row)

            headers = ["اسم الطالب", "الشعبة", "درجة الفصل الأول", "درجة نصف السنة",
                       "درجة الفصل الثاني", "درجة نهاية السنة", "مجموع الدرجات النهائية"]
            wb = Workbook()
            default_sheet = wb.active
            used_names = set()
            invalid = set('[]:*?/\\')
            for section, section_rows in groups.items():
                base = "شعبة " + section
                safe = "".join("_" if ch in invalid else ch for ch in base).strip("'")[:31] or "شعبة"
                sheet_name = safe
                suffix = 2
                while sheet_name in used_names:
                    tail = f"_{suffix}"
                    sheet_name = safe[:31-len(tail)] + tail
                    suffix += 1
                used_names.add(sheet_name)

                ws = wb.create_sheet(title=sheet_name)
                ws.sheet_view.rightToLeft = True
                ws.append(headers)
                for row in section_rows:
                    first_grade = term1_total(row)
                    second_grade = term2_total(row)
                    ws.append([row[2], row[4] or "غير محددة", excel_export_value(first_grade),
                        excel_export_value(row[10]), excel_export_value(second_grade),
                        excel_export_value(row[12]), excel_export_value(final_average(row))])

                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="182957")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                ws.row_dimensions[1].height = 34
                for col in ws.columns:
                    max_len = max(len(str(cell.value or "")) for cell in col)
                    ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 14), 32)
                    for cell in col[1:]:
                        cell.alignment = Alignment(horizontal="center", vertical="center")

            if not groups:
                default_sheet.title = "لا توجد بيانات"
                default_sheet.sheet_view.rightToLeft = True
                default_sheet.append(headers)
            else:
                wb.remove(default_sheet)
            wb.save(path)
            messagebox.showinfo("نجاح", f"تم حفظ ملف Excel واحد، ولكل شعبة ورقة مستقلة للتنقل من ألسنة الأوراق أسفل Excel:\n{path}")
        except Exception as ex:
            messagebox.showerror("خطأ", str(ex))

    def import_csv(self):
        path=filedialog.askopenfilename(filetypes=[("CSV files","*.csv")])
        if not path:return
        count, errors=0,[]
        try:
            with open(path,"r",newline="",encoding="utf-8-sig") as f:
                reader=csv.reader(f); next(reader,None)
                with db() as con:
                    for line,row in enumerate(reader,2):
                        try:
                            if len(row)<13: raise ValueError("عدد الأعمدة غير كافٍ")
                            if len(row) >= 16:
                                subject_value = row[17] if len(row) >= 18 else (row[16] if len(row) >= 17 else row[15])
                                values = (subject_value,row[1],row[2],row[3],float(row[4] or 0),float(row[5] or 0),
                                    float(row[6] or 0),float(row[7] or 0),float(row[8] or 0),float(row[10] or 0),
                                    float(row[11] or 0),float(row[13] or 0),float(row[9] or 0),float(row[12] or 0))
                                con.execute("""INSERT INTO students(subject,name,grade,section,attendance,absence,
                                    participation,activity,term1,mid_year,term2,end_year,term1_final,term2_final,
                                    attendance2,absence2,participation2,activity2)
                                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", values + tuple(values[4:8]))
                            else:
                                legacy_values = (row[1],row[2],row[3],row[4],float(row[5] or 0),float(row[6] or 0),
                                    float(row[7] or 0),float(row[8] or 0),float(row[9] or 0),
                                    float(row[10] or 0),float(row[11] or 0),float(row[12] or 0))
                                con.execute("""INSERT INTO students(subject,name,grade,section,attendance,absence,
                                    participation,activity,term1,mid_year,term2,end_year,attendance2,absence2,participation2,activity2)
                                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", legacy_values + tuple(legacy_values[4:8]))
                            count+=1
                        except Exception as ex: errors.append(f"السطر {line}: {ex}")
            self.refresh()
            msg=f"تم استيراد {count} طالبًا."
            if errors: msg+="\nأخطاء:\n"+"\n".join(errors[:8])
            messagebox.showinfo("نتيجة الاستيراد",msg)
        except Exception as ex: messagebox.showerror("خطأ",str(ex))

    @staticmethod
    def esc(value):
        s = str(value if value is not None else "")
        return s.replace("&","&amp;").replace("<","&lt;").replace(">","&gt;").replace('"',"&quot;")

    def print_report(self):
        path=filedialog.asksaveasfilename(defaultextension=".html",
            filetypes=[("HTML report","*.html")],title="حفظ تقرير A4 وفتحه للطباعة")
        if not path:return
        try:
            with db() as con: rows=con.execute("SELECT * FROM students ORDER BY section,name").fetchall()
            if not rows: messagebox.showwarning("تنبيه","لا توجد بيانات للطباعة."); return
            headers=["ت","اسم الطالب","الصف","الشعبة","حضور الفصل الأول","غياب الفصل الأول",
                "مشاركة الفصل الأول","نشاط الفصل الأول","امتحان الفصل الأول","نصف السنة",
                "حضور الفصل الثاني","غياب الفصل الثاني","مشاركة الفصل الثاني","نشاط الفصل الثاني",
                "امتحان الفصل الثاني","نهاية السنة","مجموع الفصل الأول","مجموع الفصل الثاني","المعدل النهائي"]
            body=[]
            for i,r in enumerate(rows,1):
                first_term_total = term1_total(r)
                second_term_total = term2_total(r)
                values=[i,r[2],r[3],r[4],fmt(r[5]),fmt(r[6]),fmt(r[7]),fmt(r[8]),
                        fmt(r[9]),fmt(r[10]),fmt(r[15]),fmt(r[16]),fmt(r[17]),fmt(r[18]),
                        fmt(r[11]),fmt(r[12]),fmt(first_term_total),fmt(second_term_total),fmt(final_average(r))]
                body.append("<tr>"+"".join("<td>"+self.esc(v)+"</td>" for v in values)+"</tr>")
            html=f"""<!doctype html><html lang="ar" dir="rtl"><meta charset="utf-8">
            <title>تقرير الطلاب</title><style>
            @page{{size:A4 landscape;margin:9mm}}body{{font-family:Arial;direction:rtl}}
            h2,p{{text-align:center;margin:5px}}table{{width:100%;border-collapse:collapse;font-size:9pt}}
            th,td{{border:1px solid #444;padding:5px 2px;text-align:center}}th{{background:#dce5f2}}
            </style><h2>{APP_NAME}</h2><p>{TEACHER}</p>
            <p>المادة: {self.esc(self.subject.get())} | الفصل: {self.esc(self.term.get())}</p>
            <p>تاريخ التقرير: {date.today().isoformat()}</p><table><tr>{''.join('<th>'+h+'</th>' for h in headers)}</tr>
            {''.join(body)}</table><script>window.onload=()=>window.print()</script></html>"""
            with open(path,"w",encoding="utf-8") as f:f.write(html)
            try: os.startfile(path)
            except AttributeError:
                import webbrowser; webbrowser.open("file://"+os.path.abspath(path))
            messagebox.showinfo("التقرير","تم فتح التقرير في المتصفح. اختر الطابعة من نافذة الطباعة.")
        except Exception as ex: messagebox.showerror("خطأ في الطباعة",str(ex))

    def backup(self):
        path=filedialog.asksaveasfilename(defaultextension=".db",
            filetypes=[("SQLite database","*.db")],title="حفظ نسخة احتياطية",
            initialfile="school_backup.db")
        if not path:return
        try:
            with db() as source, sqlite3.connect(path) as target: source.backup(target)
            messagebox.showinfo("نجاح",f"تم حفظ النسخة الاحتياطية:\n{path}")
        except Exception as ex:messagebox.showerror("خطأ",str(ex))

    def close(self):
        save_setting("subject",self.subject.get().strip())
        save_setting("term",self.term.get())
        self.root.destroy()

if __name__ == "__main__":
    root=tk.Tk()
    app=StudentApp(root)
    root.mainloop()
